"""Generate examples/demo-shop/knowledge/payment-policy.vi.docx (exercises the Word loader with real headings/lists)."""

from __future__ import annotations

from pathlib import Path

from docx import Document

OUT = (
    Path(__file__).resolve().parent.parent
    / "examples"
    / "demo-shop"
    / "knowledge"
    / "payment-policy.vi.docx"
)

SECTIONS: list[tuple[str, list[str], str]] = [
    (
        "1. Phương thức thanh toán được chấp nhận",
        [
            "Thanh toán khi nhận hàng (COD): trả tiền cho đơn vị vận chuyển khi nhận đơn.",
            "Chuyển khoản ngân hàng: chuyển vào tài khoản của chúng tôi, ghi mã đơn hàng ở nội dung chuyển khoản.",
            "Thẻ tín dụng và thẻ ghi nợ: Visa và Mastercard.",
            "Ví điện tử: MoMo, ZaloPay và VNPay.",
        ],
        "list",
    ),
    (
        "2. Hạn mức thanh toán khi nhận hàng",
        [
            "COD áp dụng cho đơn hàng đến 5.000.000đ. Đơn có giá trị lớn hơn phải thanh toán bằng chuyển khoản, thẻ hoặc ví điện tử."
        ],
        "para",
    ),
    (
        "3. Xác nhận thanh toán",
        [
            "Thanh toán bằng thẻ và ví điện tử được xác nhận ngay lập tức.",
            "Chuyển khoản được xác nhận trong vòng 30 phút trong giờ làm việc (08:00 đến 21:00). Chuyển khoản ngoài giờ được xác nhận vào sáng hôm sau.",
            "Đơn hàng chỉ chuyển sang trạng thái đang xử lý sau khi thanh toán được xác nhận (đơn COD chuyển sang đang xử lý ngay sau khi đặt).",
        ],
        "list",
    ),
    (
        "4. Trả góp",
        [
            "Đơn hàng từ 3.000.000đ trở lên thanh toán bằng thẻ tín dụng được hỗ trợ có thể trả góp 3 hoặc 6 tháng với lãi suất 0%. Thẻ ghi nợ và ví điện tử không áp dụng trả góp."
        ],
        "para",
    ),
    (
        "5. Hóa đơn VAT",
        [
            "Hãy yêu cầu xuất hóa đơn VAT trong vòng 24 giờ sau khi đặt hàng bằng cách cung cấp tên công ty, mã số thuế và địa chỉ. Không thể xuất hóa đơn sau khi đơn hàng đã được giao."
        ],
        "para",
    ),
    (
        "6. An toàn thanh toán",
        [
            "Chúng tôi không bao giờ yêu cầu mật khẩu thẻ, mã OTP hay số thẻ đầy đủ qua chat, điện thoại hoặc email. Nếu ai đó nhân danh chúng tôi yêu cầu các thông tin này, đừng cung cấp và hãy liên hệ hỗ trợ."
        ],
        "para",
    ),
    (
        "7. Thanh toán lỗi hoặc bị trùng",
        [
            "Nếu tiền đã bị trừ nhưng đơn hàng hiển thị chưa thanh toán, hãy chờ 30 phút rồi liên hệ hỗ trợ kèm mã đơn và ảnh chụp giao dịch. Khoản thanh toán trùng được hoàn về phương thức thanh toán ban đầu trong vòng 5 đến 7 ngày làm việc."
        ],
        "para",
    ),
]


def main() -> None:
    doc = Document()
    doc.add_heading("Chính sách thanh toán", level=1)
    for title, lines, kind in SECTIONS:
        doc.add_heading(title, level=2)
        for line in lines:
            if kind == "list":
                doc.add_paragraph(line, style="List Bullet")
            else:
                doc.add_paragraph(line)

    doc.add_heading("8. Bảng tóm tắt thời gian xác nhận", level=2)
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Phương thức"
    table.rows[0].cells[1].text = "Thời gian xác nhận"
    for method, delay in [
        ("Thẻ và ví điện tử", "Ngay lập tức"),
        ("Chuyển khoản (trong giờ làm việc)", "Trong 30 phút"),
        ("Chuyển khoản (ngoài giờ làm việc)", "Sáng hôm sau"),
        ("COD", "Khi nhận hàng"),
    ]:
        row = table.add_row().cells
        row[0].text, row[1].text = method, delay

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
