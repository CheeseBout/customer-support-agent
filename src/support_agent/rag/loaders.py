"""Document loaders for Markdown, PDF and Word. All return Markdown-flavoured text."""

from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path

SUPPORTED_SUFFIXES = {".md", ".markdown", ".pdf", ".docx"}


@dataclass(frozen=True)
class LoadedDocument:
    doc_id: str  # path relative to the knowledge dir, posix style
    source: str  # path as shown to users
    text: str  # Markdown-flavoured; headings use '#'
    content_hash: str
    updated_at: str  # ISO 8601 UTC


class UnsupportedDocument(ValueError):
    pass


def file_hash(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for block in iter(lambda: fh.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


def discover(root: Path) -> list[Path]:
    """All supported files under `root`, sorted for deterministic ingest order."""
    return sorted(
        p for p in root.rglob("*") if p.is_file() and p.suffix.lower() in SUPPORTED_SUFFIXES
    )


# --- Markdown -----------------------------------------------------------------------


def _load_markdown(path: Path) -> str:
    return path.read_text(encoding="utf-8-sig")


# --- DOCX ---------------------------------------------------------------------------


def _md_table(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    norm = [
        [c.replace("\n", " ").replace("|", "/").strip() for c in r] + [""] * (width - len(r))
        for r in rows
    ]
    lines = ["| " + " | ".join(norm[0]) + " |", "|" + "---|" * width]
    lines += ["| " + " | ".join(r) + " |" for r in norm[1:]]
    return "\n".join(lines)


def _load_docx(path: Path) -> str:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(str(path))
    out: list[str] = []
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, doc)
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name or "") if para.style is not None else ""
            if style == "Title":
                out.append(f"# {text}")
            elif m := re.match(r"Heading (\d)", style):
                out.append(f"{'#' * min(int(m.group(1)), 6)} {text}")
            elif "List" in style:
                out.append(f"- {text}")
            else:
                out.append(text)
        elif tag == "tbl":
            table = Table(child, doc)
            out.append(_md_table([[c.text for c in row.cells] for row in table.rows]))
    return "\n\n".join(out)


# --- PDF ----------------------------------------------------------------------------

_NUMBERED_HEADING = re.compile(r"^(\d+(?:\.\d+)*)[.)]?\s+\S")


def _looks_like_heading(line: str) -> bool:
    line = line.strip()
    if not line or len(line) > 90 or line.endswith((".", ",", ";", ":")):
        return False
    return bool(_NUMBERED_HEADING.match(line)) or (line.isupper() and len(line.split()) >= 2)


def _load_pdf(path: Path) -> str:
    """Text-layer PDFs only (no OCR). Numbered or ALL-CAPS short lines become headings."""
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    blocks: list[str] = []
    for page in reader.pages:
        for raw in (page.extract_text() or "").splitlines():
            line = raw.strip()
            if not line:
                blocks.append("")
            elif _looks_like_heading(line):
                blocks.append(f"\n## {line}\n")
            else:
                blocks.append(line)
    text = "\n".join(blocks)
    # Re-join hard-wrapped lines into paragraphs, keep blank lines as paragraph breaks.
    return re.sub(r"(?<!\n)\n(?!\n|#|\s*[-*|\d])", " ", text).strip()


_LOADERS = {
    ".md": _load_markdown,
    ".markdown": _load_markdown,
    ".docx": _load_docx,
    ".pdf": _load_pdf,
}


def load_document(path: Path, root: Path) -> LoadedDocument:
    suffix = path.suffix.lower()
    loader = _LOADERS.get(suffix)
    if loader is None:
        raise UnsupportedDocument(f"Unsupported file type: {suffix}")
    text = loader(path).strip()
    mtime = datetime.fromtimestamp(path.stat().st_mtime, UTC).replace(microsecond=0)
    rel = path.resolve().relative_to(root.resolve()).as_posix()
    return LoadedDocument(
        doc_id=rel,
        source=f"{root.name}/{rel}",
        text=text,
        content_hash=file_hash(path),
        updated_at=mtime.isoformat().replace("+00:00", "Z"),
    )
