"""Staff-only documents must not reach the customer-facing index."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from support_agent.core.settings import KnowledgeConfig
from support_agent.rag.audience import excluded, is_internal
from support_agent.rag.index import VectorStore
from support_agent.rag.ingest import ingest
from support_agent.rag.loaders import load_document
from tests.fakes import FakeSparse, HashingEmbeddings
from tests.test_rag import CFG, RETURN_EN, make_pdf, write

INTERNAL_MD = """# Complaint handling

**Audience:** Internal
**Version:** 1.5

## 1. Escalation

Refunds above 5,000,000 VND go to the team lead.
"""

PUBLIC_MD = """# Shipping

**Audience:** Public (customers)

## 1. Fees

Shipping inside the city costs 20,000 VND.
"""


def run(root: Path, store: VectorStore, knowledge: KnowledgeConfig | None = None):
    return ingest(
        root,
        store=store,
        embeddings=HashingEmbeddings(),
        sparse=FakeSparse(),
        embedding_model="fake",
        cfg=CFG,
        knowledge=knowledge,
    )


# --- recognising an internal document ----------------------------------------------------------------


@pytest.mark.parametrize(
    "head",
    [
        "# Plan\n\n**Audience:** Internal\n",
        "# Quy trình\n\n**Đối tượng:** Nội bộ  \n",
        "| Đối tượng | Nội bộ |\n|---|---|\n",  # a Word table row
        "Mã tài liệu COMPLAINT Đối tượng Nội bộ Phiên bản 1.3",  # text read from a PDF
        "# Handbook\n\nFor internal use only.\n",
        "# Quy định\n\nChỉ dùng nội bộ.\n",
        "# Rota\n\nStaff only.\n",
    ],
)
def test_the_head_of_a_document_can_mark_it_as_internal(head: str):
    assert is_internal(head)


@pytest.mark.parametrize(
    "text",
    [
        "# Shipping\n\n**Audience:** Public (customers)\n",
        "# Chính sách\n\n**Đối tượng:** Công khai (khách hàng)\n",
        "# Returns\n\nNothing here is about staff.\n",
        "",
    ],
)
def test_a_public_document_is_not_internal(text: str):
    assert not is_internal(text)


def test_only_the_top_of_a_document_decides():
    body = (
        "# Store policy\n\n" + "Opening hours apply.\n" * 400 + "\nThe back room is staff only.\n"
    )
    assert not is_internal(body)  # "staff only" far below the head is just a sentence


def test_a_shop_can_add_its_own_marker():
    assert not is_internal("# Memo\n\nTier: L2 only\n")
    assert is_internal("# Memo\n\nTier: L2 only\n", [r"tier:\s*l2"])


def test_paths_can_be_excluded_by_pattern():
    assert excluded("internal/rota.md", ["internal/*"])
    assert excluded("hr/staff-handbook.docx", ["*-handbook.docx"])  # matches the file name too
    assert not excluded("returns.md", ["internal/*"])


# --- ingest ------------------------------------------------------------------------------------------------


def test_an_internal_document_is_not_indexed(tmp_path: Path, store: VectorStore):
    write(tmp_path / "return.md", RETURN_EN)
    write(tmp_path / "complaints.md", INTERNAL_MD)
    report = run(tmp_path, store)
    assert report.added == ["return.md"]
    assert (
        list(report.skipped) == ["complaints.md"] and "internal" in report.skipped["complaints.md"]
    )
    assert set(store.indexed_documents()) == {"return.md"}
    assert "skipped=1" in report.summary()


def test_a_public_document_with_an_audience_line_is_indexed(tmp_path: Path, store: VectorStore):
    write(tmp_path / "ship.md", PUBLIC_MD)
    assert run(tmp_path, store).added == ["ship.md"]


def test_a_document_indexed_earlier_is_removed_once_it_is_known_to_be_internal(
    tmp_path: Path, store: VectorStore
):
    path = write(tmp_path / "plan.md", PUBLIC_MD)
    run(tmp_path, store)
    assert "plan.md" in store.indexed_documents()
    path.write_text(INTERNAL_MD, encoding="utf-8")  # the shop marks it internal later
    report = run(tmp_path, store)
    assert "plan.md" in report.skipped and report.removed == ["plan.md"]
    assert store.indexed_documents() == {}


def test_a_document_is_removed_when_the_shop_starts_excluding_its_path(
    tmp_path: Path, store: VectorStore
):
    write(tmp_path / "private" / "notes.md", PUBLIC_MD)
    write(tmp_path / "return.md", RETURN_EN)
    run(tmp_path, store)
    report = run(tmp_path, store, KnowledgeConfig(exclude=["private/*"]))
    assert report.removed == ["private/notes.md"] and list(report.skipped) == ["private/notes.md"]
    assert set(store.indexed_documents()) == {"return.md"}


def test_the_check_can_be_switched_off(tmp_path: Path, store: VectorStore):
    write(tmp_path / "complaints.md", INTERNAL_MD)
    report = run(tmp_path, store, KnowledgeConfig(skip_internal=False))
    assert report.added == ["complaints.md"] and report.skipped == {}


def test_word_and_pdf_documents_are_recognised_too(tmp_path: Path, store: VectorStore):
    docx = Document()
    docx.add_heading("Quy trình xử lý khiếu nại", 1)
    table = docx.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Đối tượng"
    table.rows[0].cells[1].text = "Nội bộ"
    docx.add_paragraph("Hoàn tiền trên 5.000.000đ do trưởng nhóm duyệt.")
    docx.save(str(tmp_path / "process.docx"))
    make_pdf(tmp_path / "rota.pdf", ["Staff rota", "Audience Internal", "Morning shift 8:30"])
    assert is_internal(load_document(tmp_path / "process.docx", tmp_path).text)
    assert is_internal(load_document(tmp_path / "rota.pdf", tmp_path).text)
    report = run(tmp_path, store)
    assert sorted(report.skipped) == ["process.docx", "rota.pdf"] and report.added == []
